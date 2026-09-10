import openpyxl
import json
import re

EXCEL_PATH = r'E:\微信聊天记录导出\Mia\20260907\微信聊天记录_wxid_hm4h93c1b3n122\需求联系人分析报告_新.xlsx'
OUTPUT_PATH = r'E:\一人公司\public\data.json'

REGION_KEYWORDS = [
    '悉尼', '墨尔本', '布里斯班', '阿德莱德', '珀斯', '堪培拉', '新西兰', '奥克兰', '惠灵顿',
    '迪拜', '阿联酋', '新加坡', '日本', '东京', '大阪', '英国', '伦敦', '美国', '加拿大', '温哥华',
    '多伦多', 'city', 'CBD', 'Ryde', 'Box Hill', 'Burwood', 'Hurstville', 'Chatswood',
    'Parramatta', 'Strathfield', 'Bankstown', 'Fairfield', 'Liverpool', 'Campbelltown',
    'Blacktown', 'Penrith', 'Newcastle', 'Wollongong', 'Geelong', '黄金海岸', '拉斯维加斯',
    '达拉斯', '旧金山', '洛杉矶', '纽约', '芝加哥', '休斯顿', '凤凰城', '西雅图',
]

COMPETITOR_KEYWORDS = [
    '广告', 'Google', '刷好评', '删差评', '投流', '排名', '代发', '换群',
    '诚招代理', '招代理', '源头价格', 'SKU', 'xs价格', '市场比价',
    '货代', '船公司', '转运', '庄家', '一代', '工签',
    '签证', '移民', '建群', '拉你入群', '商家网站', '好评',
    '代理', '出方案', '投流置顶', '店铺维护', 'Facebook',
]

JOB_TYPE_KEYWORDS = {
    '厨师/餐饮': ['厨师', '中餐', '川菜', '粤菜', '湘菜', '西餐', '餐饮', '炒菜', '大锅饭', '小炒', '食堂', '料理', '拉面', '寿司'],
    '按摩/SPA': ['按摩', 'SPA', '足疗', '刮痧', '采耳', '技师', '足浴', '推拿', '精油'],
    '装修/建筑': ['装修', '建筑', '油漆', '木工', '泥工', '电工', '水管', '地板', '瓷砖', '大工', '中工', '小工', '工地', '铁工', 'plasterboard', '做铁', '板灰', '砌砖'],
    '美容/美甲': ['美容', '美甲', '美睫', '纹绣', '化妆', '皮肤管理', '美发', '理发'],
    '司机/物流': ['司机', '代驾', '送货', '物流', '叉车', '面包车', '驾照', 'Uber', 'Eats'],
    '保姆/护工': ['保姆', '月嫂', '护工', '陪护', '清洁', '保洁', '住家', '照顾', '老人'],
    '教育/培训': ['教育', '培训', '老师', '中文', '英语', '辅导', '幼教', '学校'],
    '零售/销售': ['销售', '店长', '收银', '导购', '零售', '超市', '店铺', '门店'],
    '工厂/制造': ['工厂', '制造', '仓库', '包装', '分拣', '流水线', '普工'],
    '其他': [],
}

def clean_text(text):
    text = re.sub(r'\[庆祝\]|\[烟花\]|\[爱心\]|\[红包\]|\[握手\]|\[玫瑰\]|\[發\]|\[福\]', '', text)
    text = re.sub(r'[\U0001f389-\U0001f9ff\U0001fa70-\U0001faff\U00002702-\U000027b0\U0001f300-\U0001f5ff\U0001f600-\U0001f64f\U0001f680-\U0001f6ff\U0001f900-\U0001f9ff\U0001fa00-\U0001fa6f\U0001fa70-\U0001faff]', '', text)
    text = re.sub(r'\d️⃣|\u200d', '', text)
    text = re.sub(r'\[[\u4e00-\u9fa5]{1,6}\]', '', text)
    text = re.sub(r'\n{3,}', '\n\n', text)
    return text.strip()

def extract_region(text):
    for kw in REGION_KEYWORDS:
        if kw.lower() in text.lower():
            return kw
    return ''

def extract_job_type(desc):
    desc_lower = desc.lower()
    for job_type, keywords in JOB_TYPE_KEYWORDS.items():
        for kw in keywords:
            if kw.lower() in desc_lower:
                return job_type
    return '其他'

def extract_title(desc, nickname):
    lines = [l.strip() for l in desc.split('\n') if l.strip()]
    title = ''
    for line in lines[:8]:
        line = clean_text(line)
        if not line or len(line) < 4:
            continue
        if any(kw in line for kw in ['招', '聘', '急招', '诚聘', '需要', '找人', '诚招', '招工']):
            title = line
            break
    if not title and lines:
        for line in lines[:5]:
            cleaned = clean_text(line)
            if cleaned and len(cleaned) >= 4:
                title = cleaned
                break
    if len(title) > 30:
        title = title[:30] + '...'
    return title or (nickname[:20] if nickname else '招聘信息')

def is_competitor(desc, nickname):
    text = (desc + ' ' + str(nickname)).lower()
    for kw in COMPETITOR_KEYWORDS:
        if kw.lower() in text:
            return True
    return False

wb = openpyxl.load_workbook(EXCEL_PATH)
ws = wb.active

records = []
filtered_count = 0
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
    category = str(row[7] or '')
    if '招聘求职' not in category:
        continue

    desc = (row[8] or '').strip()
    nickname = row[3] or row[2] or ''

    if is_competitor(desc, nickname):
        filtered_count += 1
        continue

    region = row[4] or extract_region(desc)
    title = extract_title(desc, nickname)
    description = clean_text(desc)
    phone = str(row[9] or '')
    other_contact = str(row[10] or '')
    publish_date = str(row[11] or '')[:10]
    job_type = extract_job_type(desc)

    record = {
        'id': row[0],
        'title': title,
        'nickname': nickname,
        'region': region,
        'description': description,
        'phone': phone,
        'other_contact': other_contact,
        'tags': [t.strip() for t in category.split('、') if t.strip()],
        'job_type': job_type,
        'publish_date': publish_date,
        'contact_name': 'Mia',
    }
    records.append(record)

print(f"Total: {len(records)}, Filtered: {filtered_count}")

with open(OUTPUT_PATH, 'w', encoding='utf-8') as f:
    json.dump(records, f, ensure_ascii=False, indent=2)

print(f"Saved to {OUTPUT_PATH}")
